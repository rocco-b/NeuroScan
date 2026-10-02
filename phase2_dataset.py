"""Phase-2 dataset builder for NeuroScan.

This module turns a folder of EMG/NCS Word reports into a de-identified
validation package: anonymized DOCX copies, nerve-level measurements,
NeuroScan predictions, QC reports, an expert-label template and a SQLite
archive.  It is deliberately deterministic and audit-friendly.
"""
from __future__ import annotations

import argparse
import csv
import datetime as _dt
import os
import re
import shutil
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, Iterator, List, Optional, Tuple

from docx import Document

from analysis_logic import analyze_data, validate_exam
from file_processor import FileProcessor

SUPPORTED_SUFFIXES = {".doc", ".docx"}

FIELD_LABELS = [
    "PATIENT", "ADDRESS STREET", "ADDRESS", "CITY", "INSURANCE", "ID NUMBER",
    "PHYSICIAN", "REF PHYS", "TECHNICIAN", "DATE OF BIRTH", "DOB",
]
DATE_FIELD_LABELS = ["EMG EXAMINATION DATE", "DATE OF REPORT"]
NEXT_LABELS = FIELD_LABELS + DATE_FIELD_LABELS + ["AGE", "HEIGHT", "DIAGNOSIS", "MEDICAL HISTORY"]

ITALIAN_FISCAL_CODE_RE = re.compile(r"\b[A-Z]{6}\d{2}[A-Z]\d{2}[A-Z]\d{3}[A-Z]\b")
DATE_RE = re.compile(r"\b\d{1,2}[/-]\d{1,2}[/-](?:\d{2}|\d{4})(?:\s+\d{1,2}:\d{2}(?::\d{2})?)?\b")
EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I)
PHONE_RE = re.compile(r"(?<!\d)(?:\+?39\s*)?(?:0\d{1,4}[\s./-]?\d{5,8}|3\d{2}[\s./-]?\d{6,7})(?!\d)")
SIGNATURE_RE = re.compile(
    r"(?i)\b(?:Prof\.?|Dott\.?|Dr\.?)\s+"
    r"(?:[A-ZÀ-ÖØ-Ý](?:\.)?\s*){0,3}"
    r"[A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÖØ-öø-ÿ'.-]+"
    r"(?:\s+[A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÖØ-öø-ÿ'.-]+){0,3}"
)

_FIELD_VALUE_RE = re.compile(
    rf"(?i)\b({'|'.join(re.escape(x) for x in FIELD_LABELS)})\s*:\s*"
    rf".*?(?=\s+\b(?:{'|'.join(re.escape(x) for x in NEXT_LABELS)})\s*:|$)"
)
_DATE_FIELD_RE = re.compile(
    rf"(?i)\b({'|'.join(re.escape(x) for x in DATE_FIELD_LABELS)})\s*:\s*"
    rf".*?(?=\s+\b(?:{'|'.join(re.escape(x) for x in NEXT_LABELS)})\s*:|$)"
)


@dataclass
class AnonymizationResult:
    case_id: str
    input_file: str
    output_file: str
    replacements: Dict[str, int]
    residual_findings: Dict[str, int]
    status: str
    message: str = ""


def _iter_story_containers(document: Document) -> Iterator[object]:
    """Yield document body plus unique header/footer story containers."""
    yield document
    seen = set()
    for section in document.sections:
        for container in (
            section.header, section.first_page_header, section.even_page_header,
            section.footer, section.first_page_footer, section.even_page_footer,
        ):
            key = id(container._element)
            if key not in seen:
                seen.add(key)
                yield container


def _iter_all_tables(container) -> Iterator[object]:
    """Yield top-level and nested tables from a document story container."""
    for table in getattr(container, "tables", []):
        yield table
        for row in table.rows:
            seen_cells = set()
            for cell in row.cells:
                key = id(cell._tc)
                if key in seen_cells:
                    continue
                seen_cells.add(key)
                yield from _iter_all_tables(cell)


def _iter_docx_paragraphs(document: Document):
    """Yield body/header/footer paragraphs, including paragraphs in tables."""
    seen = set()
    for container in _iter_story_containers(document):
        for paragraph in getattr(container, "paragraphs", []):
            key = id(paragraph._p)
            if key not in seen:
                seen.add(key)
                yield paragraph
        for table in _iter_all_tables(container):
            for row in table.rows:
                seen_cells = set()
                for cell in row.cells:
                    cell_key = id(cell._tc)
                    if cell_key in seen_cells:
                        continue
                    seen_cells.add(cell_key)
                    for paragraph in cell.paragraphs:
                        key = id(paragraph._p)
                        if key not in seen:
                            seen.add(key)
                            yield paragraph


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    chunks: List[str] = []
    for paragraph in _iter_docx_paragraphs(doc):
        if paragraph.text:
            chunks.append(paragraph.text)
    return "\n".join(chunks)


def extract_report_metadata(docx_path: str | Path) -> Dict[str, Optional[float | int | str]]:
    """Extract non-identifying metadata needed for analysis."""
    text = _docx_text(Path(docx_path))
    metadata: Dict[str, Optional[float | int | str]] = {
        "age_years": None,
        "height_cm": None,
        "exam_year": None,
        "report_year": None,
    }

    age_match = re.search(r"\bAGE\s*:\s*(\d{1,3})\b", text, re.I)
    if age_match:
        metadata["age_years"] = int(age_match.group(1))

    height_match = re.search(r"\bHEIGHT\s*:\s*(\d{2,3}(?:[,.]\d+)?)", text, re.I)
    if height_match:
        metadata["height_cm"] = float(height_match.group(1).replace(",", "."))

    exam_date = re.search(r"\bEMG EXAMINATION DATE\s*:\s*(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", text, re.I)
    if exam_date:
        year = int(exam_date.group(3))
        metadata["exam_year"] = year + 2000 if year < 100 else year

    report_date = re.search(r"\bDATE OF REPORT\s*:\s*(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", text, re.I)
    if report_date:
        year = int(report_date.group(3))
        metadata["report_year"] = year + 2000 if year < 100 else year

    return metadata


def sanitize_text(text: str, redact_dates: bool = True) -> Tuple[str, Dict[str, int]]:
    """Return sanitized text plus replacement counts by category."""
    if not text:
        return text, {}

    counts: Dict[str, int] = {}

    def count_sub(pattern: re.Pattern, repl, category: str, value: str) -> str:
        new_value, n = pattern.subn(repl, value)
        if n:
            counts[category] = counts.get(category, 0) + n
        return new_value

    def repl_field(match: re.Match) -> str:
        label = match.group(1).upper()
        return f"{label}: [REDACTED]"

    def repl_date_field(match: re.Match) -> str:
        label = match.group(1).upper()
        return f"{label}: [REDACTED_DATE]"

    text = count_sub(_FIELD_VALUE_RE, repl_field, "field_values", text)
    if redact_dates:
        text = count_sub(_DATE_FIELD_RE, repl_date_field, "date_fields", text)
        text = count_sub(DATE_RE, "[REDACTED_DATE]", "dates", text)
    text = count_sub(ITALIAN_FISCAL_CODE_RE, "[REDACTED_ID]", "italian_fiscal_codes", text)
    text = count_sub(EMAIL_RE, "[REDACTED_EMAIL]", "emails", text)
    text = count_sub(PHONE_RE, "[REDACTED_PHONE]", "phones", text)
    text = count_sub(SIGNATURE_RE, "[REDACTED_SIGNATURE]", "signatures", text)
    return text, counts


def scan_sensitive_text(text: str) -> Dict[str, int]:
    """Count common residual identifiers in already-sanitized text."""
    findings = {
        "italian_fiscal_codes": len(ITALIAN_FISCAL_CODE_RE.findall(text or "")),
        "dates": len(DATE_RE.findall(text or "")),
        "emails": len(EMAIL_RE.findall(text or "")),
        "phones": len(PHONE_RE.findall(text or "")),
        "signatures": len(SIGNATURE_RE.findall(text or "")),
    }
    for label in FIELD_LABELS:
        # Flag labels only when the same line appears to still contain unredacted text.
        pattern = re.compile(rf"(?i)\b{re.escape(label)}\s*:\s*(?!\[REDACTED(?:_[A-Z]+)?\])\S+")
        findings[f"unredacted_{label.lower().replace(' ', '_')}"] = len(pattern.findall(text or ""))
    return {k: v for k, v in findings.items() if v}


_ALL_LABELS = sorted(set(NEXT_LABELS), key=len, reverse=True)
_SENSITIVE_LABEL_SET = {x.upper() for x in FIELD_LABELS}
_DATE_LABEL_SET = {x.upper() for x in DATE_FIELD_LABELS}
_LABEL_AT_START_RE = re.compile(
    rf"(?i)^\s*({'|'.join(re.escape(x) for x in _ALL_LABELS)})\s*:"
)


def _label_at_start(text: str) -> Optional[str]:
    match = _LABEL_AT_START_RE.search(text or "")
    return match.group(1).upper() if match else None


def _unique_cells(row) -> List[object]:
    cells = []
    seen = set()
    for cell in row.cells:
        key = id(cell._tc)
        if key not in seen:
            seen.add(key)
            cells.append(cell)
    return cells


def _set_cell_text(cell, text: str) -> None:
    """Replace a cell value deterministically, including duplicated runs."""
    cell.text = text


def _redact_table_structurally(table) -> Dict[str, int]:
    """Redact values stored in cells adjacent to administrative labels.

    DY11/LibreOffice reports often place ``PATIENT:`` and the patient value in
    separate cells.  Paragraph-only regex replacement cannot remove the latter.
    This routine uses row structure and also handles merged/duplicated cells.
    """
    counts: Dict[str, int] = {}
    for row in table.rows:
        cells = _unique_cells(row)
        labels = [_label_at_start(cell.text) for cell in cells]
        for i, label in enumerate(labels):
            if label not in _SENSITIVE_LABEL_SET and label not in _DATE_LABEL_SET:
                continue
            replacement = f"{label}: [REDACTED_DATE]" if label in _DATE_LABEL_SET or label in {"DATE OF BIRTH", "DOB"} else f"{label}: [REDACTED]"
            if cells[i].text.strip() != replacement:
                _set_cell_text(cells[i], replacement)
                counts["label_cells"] = counts.get("label_cells", 0) + 1

            # Clear every value cell until the next recognized field label.
            j = i + 1
            while j < len(cells) and labels[j] is None:
                if cells[j].text.strip():
                    _set_cell_text(cells[j], "")
                    counts["adjacent_value_cells"] = counts.get("adjacent_value_cells", 0) + 1
                j += 1
    return counts


def _scan_structural_residuals(document: Document) -> Dict[str, int]:
    """Detect residual values adjacent to sensitive labels after redaction."""
    findings: Dict[str, int] = {}
    for container in _iter_story_containers(document):
        for table in _iter_all_tables(container):
            for row in table.rows:
                cells = _unique_cells(row)
                labels = [_label_at_start(cell.text) for cell in cells]
                for i, label in enumerate(labels):
                    if label not in _SENSITIVE_LABEL_SET and label not in _DATE_LABEL_SET:
                        continue
                    j = i + 1
                    while j < len(cells) and labels[j] is None:
                        value = cells[j].text.strip()
                        if value and not value.startswith("[REDACTED"):
                            key = f"adjacent_{label.lower().replace(' ', '_')}"
                            findings[key] = findings.get(key, 0) + 1
                        j += 1
    return findings


def scan_docx_sensitive(path: str | Path) -> Dict[str, int]:
    """Combined textual and structural post-anonymization privacy scan."""
    path = Path(path)
    doc = Document(str(path))
    findings = scan_sensitive_text(_docx_text(path))
    for key, value in _scan_structural_residuals(doc).items():
        findings[key] = findings.get(key, 0) + value
    return {k: v for k, v in findings.items() if v}


def anonymize_docx(input_path: str | Path, output_path: str | Path, case_id: str) -> AnonymizationResult:
    """Create an anonymized DOCX copy while preserving analyzable NCS tables.

    Redaction is performed both structurally (label/value cells) and textually.
    A combined post-save scan is mandatory; a document passes only when no
    supported residual identifier pattern remains.
    """
    input_path = Path(input_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    replacements: Dict[str, int] = {"case_id_inserted": 1}

    try:
        doc = Document(str(input_path))

        # First remove values in cells adjacent to administrative labels.
        for container in _iter_story_containers(doc):
            for table in _iter_all_tables(container):
                counts = _redact_table_structurally(table)
                for key, value in counts.items():
                    replacements[key] = replacements.get(key, 0) + value

        # Then sanitize residual inline values, free text, headers and footers.
        for paragraph in _iter_docx_paragraphs(doc):
            cleaned, counts = sanitize_text(paragraph.text)
            if cleaned != paragraph.text:
                paragraph.text = cleaned
            for key, value in counts.items():
                replacements[key] = replacements.get(key, 0) + value

        # Avoid inserting duplicate IDs when re-scrubbing an earlier output.
        existing_case_id = any(
            p.text.strip().upper().startswith("NEUROSCAN_CASE_ID:")
            for p in doc.paragraphs[:5]
        )
        if not existing_case_id:
            first = doc.paragraphs[0] if doc.paragraphs else None
            if first is not None:
                first.insert_paragraph_before(f"NEUROSCAN_CASE_ID: {case_id}")
            else:
                doc.add_paragraph(f"NEUROSCAN_CASE_ID: {case_id}")
        else:
            replacements["case_id_inserted"] = 0

        doc.save(str(output_path))
        residual = scan_docx_sensitive(output_path)
        status = "PASS" if not residual else "REVIEW_REQUIRED"
        return AnonymizationResult(
            case_id=case_id,
            input_file=str(input_path),
            output_file=str(output_path),
            replacements=replacements,
            residual_findings=residual,
            status=status,
            message="" if status == "PASS" else "Residual potentially identifying patterns were detected by the post-save scan.",
        )
    except Exception as exc:
        return AnonymizationResult(
            case_id=case_id,
            input_file=str(input_path),
            output_file=str(output_path),
            replacements=replacements,
            residual_findings={},
            status="ERROR",
            message=str(exc),
        )


def _write_csv(path: Path, rows: List[Dict], fieldnames: List[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _append_count_dict(prefix: str, values: Dict[str, int]) -> str:
    return "; ".join(f"{k}={v}" for k, v in sorted(values.items())) if values else ""


def _flatten_measurements(case_id: str, data: Dict[str, List[dict]]) -> List[Dict]:
    rows: List[Dict] = []
    for modality, records in data.items():
        for idx, rec in enumerate(records, start=1):
            rows.append({
                "case_id": case_id,
                "modality": modality,
                "row_index": idx,
                "nerve": rec.get("Nerve", ""),
                "amp": rec.get("Amp", rec.get("M-Amp", "")),
                "cv": rec.get("CV", ""),
                "dml": rec.get("DML", ""),
                "fmin": rec.get("Fmin", ""),
                "persist": rec.get("Persist", ""),
            })
    return rows


def _flatten_predictions(case_id: str, results: Dict) -> List[Dict]:
    rows: List[Dict] = []
    pnp = results.get("pnp_classification", {}) or {}
    for modality in ("motor", "sensory", "fwave"):
        for rec in results.get(modality, []):
            ev = rec.get("evidence", {}) if isinstance(rec, dict) else {}
            z = ev.get("z_scores", {}) if isinstance(ev, dict) else {}
            rows.append({
                "case_id": case_id,
                "modality": modality,
                "nerve": rec.get("nerve", ""),
                "nerve_diagnosis": rec.get("diagnosis", ""),
                "patient_pnp_pattern": pnp.get("pattern", ""),
                "patient_pnp_details": pnp.get("details", ""),
                "z_amp": z.get("Amp", ""),
                "z_cv": z.get("CV", ""),
                "z_dml": z.get("DML", ""),
                "z_fmin": z.get("Fmin", ""),
                "custom_refs": ev.get("reference_override_applied", ""),
            })
    return rows


def _ensure_sqlite(output_path: Path, csv_paths: Dict[str, Path]) -> None:
    if output_path.exists():
        output_path.unlink()
    conn = sqlite3.connect(str(output_path))
    try:
        for table_name, csv_path in csv_paths.items():
            with csv_path.open("r", newline="", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                fieldnames = reader.fieldnames or []
                if not fieldnames:
                    continue
                cols = ", ".join(f'"{c}" TEXT' for c in fieldnames)
                conn.execute(f'CREATE TABLE "{table_name}" ({cols})')
                placeholders = ", ".join("?" for _ in fieldnames)
                colnames = ", ".join(f'"{c}"' for c in fieldnames)
                for row in reader:
                    conn.execute(
                        f'INSERT INTO "{table_name}" ({colnames}) VALUES ({placeholders})',
                        [str(row.get(c, "")) for c in fieldnames],
                    )
        conn.commit()
    finally:
        conn.close()


def discover_word_files(input_dir: str | Path) -> List[Path]:
    """Discover both modern .docx and legacy binary .doc reports."""
    input_dir = Path(input_dir)
    files: List[Path] = []
    for suffix in SUPPORTED_SUFFIXES:
        files.extend(input_dir.rglob(f"*{suffix}"))
    return sorted(
        p for p in files
        if not p.name.startswith("~$")
        and p.suffix.lower() in SUPPORTED_SUFFIXES
    )


def discover_docx_files(input_dir: str | Path) -> List[Path]:
    """Backward-compatible alias retained for older tests/scripts."""
    return discover_word_files(input_dir)


def prepare_docx_for_processing(file_path: str | Path, conversion_dir: str | Path, *, doc_converter: str = "auto", conversion_timeout: int = 60) -> Path:
    """Return a .docx path suitable for python-docx parsing.

    .docx files are returned unchanged. Legacy .doc files are converted into
    conversion_dir, which must be treated as sensitive because the intermediate
    document still contains original identifiers.
    """
    file_path = Path(file_path)
    suffix = file_path.suffix.lower()
    if suffix == ".docx":
        return file_path
    if suffix == ".doc":
        conversion_dir = Path(conversion_dir)
        conversion_dir.mkdir(parents=True, exist_ok=True)
        return Path(FileProcessor(str(file_path))._convert_doc_to_docx(str(file_path), output_dir=str(conversion_dir), converter=doc_converter, timeout=conversion_timeout))
    raise ValueError(f"Unsupported report format: {file_path.suffix}")


def build_validation_dataset(
    input_dir: str | Path,
    output_dir: str | Path,
    *,
    age_default: float = 40,
    height_default: float = 170,
    anonymize: bool = True,
    case_prefix: str = "NS",
    doc_converter: str = "auto",
    conversion_timeout: int = 60,
) -> Dict[str, object]:
    """Build the full Phase-2 validation package from a folder of .doc/.docx reports."""
    input_dir = Path(input_dir)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    anon_dir = output_dir / "anonymized_docx"
    private_dir = output_dir / "private_mapping_STORE_SEPARATELY"
    private_dir.mkdir(parents=True, exist_ok=True)

    conversion_dir = private_dir / "converted_docx_sensitive_DO_NOT_SHARE"
    discovered_files = discover_word_files(input_dir)
    files = [
        p for p in discovered_files
        if output_dir.resolve() not in p.resolve().parents
        and "anonymized_docx" not in p.parts
        and "private_mapping_STORE_SEPARATELY" not in p.parts
    ]
    case_index_rows: List[Dict] = []
    anonymization_rows: List[Dict] = []
    parsing_rows: List[Dict] = []
    measurement_rows: List[Dict] = []
    prediction_rows: List[Dict] = []
    expert_template_rows: List[Dict] = []
    private_mapping_rows: List[Dict] = []

    used_case_ids = set()
    for idx, file_path in enumerate(files, start=1):
        existing_id = re.fullmatch(r"(?i)(NS\d{5})", file_path.stem.strip())
        case_id = existing_id.group(1).upper() if existing_id else f"{case_prefix}{idx:05d}"
        if case_id in used_case_ids:
            raise ValueError(f"Duplicate case_id derived from input files: {case_id}")
        used_case_ids.add(case_id)
        anon_path = anon_dir / f"{case_id}.docx"
        metadata: Dict[str, Optional[float | int | str]] = {
            "age_years": None, "height_cm": None, "exam_year": None, "report_year": None
        }
        analysis_docx: Optional[Path] = None

        try:
            analysis_docx = prepare_docx_for_processing(file_path, conversion_dir, doc_converter=doc_converter, conversion_timeout=conversion_timeout)
            metadata = extract_report_metadata(analysis_docx)
        except Exception as exc:
            message = f"Conversion/metadata extraction failed: {exc}"
            anonymization_rows.append({
                "case_id": case_id,
                "input_file_name": f"{case_id}.docx",
                "anonymized_file": "",
                "status": "ERROR",
                "message": message,
                "replacement_counts": "",
                "residual_findings": "",
            })
            parsing_rows.append({
                "case_id": case_id,
                "file_name": f"{case_id}.docx",
                "status": "ERROR",
                "motor_rows": 0,
                "sensory_rows": 0,
                "fwave_rows": 0,
                "exam_completeness_status": "",
                "exam_completeness_issues": "",
                "message": message,
            })
            expert_template_rows.append({
                "case_id": case_id,
                "software_pnp_pattern": "CONVERSION_ERROR",
                "expert_1_label": "",
                "expert_2_label": "",
                "consensus_label": "",
                "review_notes": message,
            })
            case_index_rows.append({
                "case_id": case_id,
                "age_years": "",
                "height_cm": "",
                "exam_year": "",
                "report_year": "",
                "source_file_name": f"{case_id}.docx",
            })
            private_mapping_rows.append({
                "case_id": case_id,
                "original_path": str(file_path.resolve()),
                "original_name": file_path.name,
                "anonymized_name": "",
                "converted_docx_sensitive": "",
            })
            continue

        age = float(metadata.get("age_years") or age_default)
        height = float(metadata.get("height_cm") or height_default)

        if anonymize:
            anon_result = anonymize_docx(analysis_docx, anon_path, case_id)
        else:
            anon_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(analysis_docx, anon_path)
            anon_result = AnonymizationResult(case_id, str(file_path), str(anon_path), {}, {}, "NOT_RUN", "")

        anonymization_rows.append({
            "case_id": case_id,
            "input_file_name": anon_path.name,
            "anonymized_file": str(anon_path.relative_to(output_dir)) if anon_path.exists() else "",
            "status": anon_result.status,
            "message": anon_result.message,
            "replacement_counts": _append_count_dict("", anon_result.replacements),
            "residual_findings": _append_count_dict("", anon_result.residual_findings),
        })
        private_mapping_rows.append({
            "case_id": case_id,
            "original_path": str(file_path.resolve()),
            "original_name": file_path.name,
            "anonymized_name": anon_path.name,
            "converted_docx_sensitive": str(analysis_docx.resolve()) if analysis_docx and analysis_docx != file_path else "",
        })

        try:
            parse_path = anon_path if anonymize and anon_path.exists() else analysis_docx
            processor = FileProcessor(str(parse_path))
            data = processor._extract_data(str(parse_path))
            parser_errors = list(processor.last_diagnostics.get("parser_errors", []))
            if parser_errors:
                raise RuntimeError("; ".join(parser_errors))
            validation = validate_exam(data.get("motor", []), data.get("sensory", []), data.get("fwave", []))
            results = analyze_data(
                data.get("motor", []), data.get("sensory", []), data.get("fwave", []),
                patient_age=age, patient_height=height,
            )
            parsing_status = "PASS" if any(data.get(k) for k in ("motor", "sensory", "fwave")) else "REVIEW_REQUIRED"
            if anon_result.status != "PASS":
                parsing_status = "REVIEW_REQUIRED"
            measurement_rows.extend(_flatten_measurements(case_id, data))
            prediction_rows.extend(_flatten_predictions(case_id, results))
            pnp = results.get("pnp_classification", {}) or {}
            expert_template_rows.append({
                "case_id": case_id,
                "software_pnp_pattern": pnp.get("pattern", ""),
                "expert_1_label": "",
                "expert_2_label": "",
                "consensus_label": "",
                "review_notes": "",
            })
            parsing_rows.append({
                "case_id": case_id,
                "file_name": f"{case_id}.docx",
                "status": parsing_status,
                "motor_rows": len(data.get("motor", [])),
                "sensory_rows": len(data.get("sensory", [])),
                "fwave_rows": len(data.get("fwave", [])),
                "exam_completeness_status": "PASS" if validation.get("is_valid", False) else "REVIEW_REQUIRED",
                "exam_completeness_issues": "; ".join((validation.get("errors", []) or []) + (validation.get("warnings", []) or [])) if isinstance(validation, dict) else "",
                "message": "",
            })
        except Exception as exc:
            parsing_rows.append({
                "case_id": case_id,
                "file_name": f"{case_id}.docx",
                "status": "ERROR",
                "motor_rows": 0,
                "sensory_rows": 0,
                "fwave_rows": 0,
                "exam_completeness_status": "",
                "exam_completeness_issues": "",
                "message": str(exc),
            })
            expert_template_rows.append({
                "case_id": case_id,
                "software_pnp_pattern": "PARSING_ERROR",
                "expert_1_label": "",
                "expert_2_label": "",
                "consensus_label": "",
                "review_notes": str(exc),
            })

        case_index_rows.append({
            "case_id": case_id,
            "age_years": metadata.get("age_years") or "",
            "height_cm": metadata.get("height_cm") or "",
            "exam_year": metadata.get("exam_year") or "",
            "report_year": metadata.get("report_year") or "",
            "source_file_name": anon_path.name,
        })

    reports_dir = output_dir / "reports"
    csv_paths = {
        "case_index": reports_dir / "case_index.csv",
        "anonymization_report": reports_dir / "anonymization_report.csv",
        "parsing_report": reports_dir / "parsing_report.csv",
        "nerve_measurements": reports_dir / "nerve_measurements.csv",
        "neuroscan_predictions": reports_dir / "neuroscan_predictions.csv",
        "expert_labels_template": reports_dir / "expert_labels_template.csv",
    }

    _write_csv(csv_paths["case_index"], case_index_rows, ["case_id", "age_years", "height_cm", "exam_year", "report_year", "source_file_name"])
    _write_csv(csv_paths["anonymization_report"], anonymization_rows, ["case_id", "input_file_name", "anonymized_file", "status", "message", "replacement_counts", "residual_findings"])
    _write_csv(csv_paths["parsing_report"], parsing_rows, ["case_id", "file_name", "status", "motor_rows", "sensory_rows", "fwave_rows", "exam_completeness_status", "exam_completeness_issues", "message"])
    _write_csv(csv_paths["nerve_measurements"], measurement_rows, ["case_id", "modality", "row_index", "nerve", "amp", "cv", "dml", "fmin", "persist"])
    _write_csv(csv_paths["neuroscan_predictions"], prediction_rows, ["case_id", "modality", "nerve", "nerve_diagnosis", "patient_pnp_pattern", "patient_pnp_details", "z_amp", "z_cv", "z_dml", "z_fmin", "custom_refs"])
    _write_csv(csv_paths["expert_labels_template"], expert_template_rows, ["case_id", "software_pnp_pattern", "expert_1_label", "expert_2_label", "consensus_label", "review_notes"])
    _write_csv(private_dir / "case_mapping_private.csv", private_mapping_rows, ["case_id", "original_path", "original_name", "anonymized_name", "converted_docx_sensitive"])

    sqlite_path = output_dir / "neuroscan_validation.sqlite"
    _ensure_sqlite(sqlite_path, csv_paths)

    readme_path = output_dir / "README_PHASE2_OUTPUT.txt"
    readme_path.write_text(
        "NeuroScan Phase-2 validation package\n"
        f"Created: {_dt.datetime.now().isoformat(timespec='seconds')}\n"
        f"Input folder: {input_dir}\n"
        f"Number of Word files processed (.doc/.docx): {len(files)}\n\n"
        "Use reports/expert_labels_template.csv for blinded expert review.\n"
        "The private_mapping_STORE_SEPARATELY folder contains direct source-file mapping and must not be shared.\n"
        "Files with REVIEW_REQUIRED or ERROR status require manual inspection before publication.\n",
        encoding="utf-8",
    )

    return {
        "input_dir": str(input_dir),
        "output_dir": str(output_dir),
        "files_processed": len(files),
        "reports_dir": str(reports_dir),
        "sqlite_path": str(sqlite_path),
        "review_required": sum(1 for r in anonymization_rows + parsing_rows if r.get("status") in {"REVIEW_REQUIRED", "ERROR"}),
    }


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Build a NeuroScan Phase-2 validation dataset from .doc/.docx reports.")
    parser.add_argument("--input", required=True, help="Folder containing original .doc/.docx reports")
    parser.add_argument("--output", required=True, help="Output folder for the validation package")
    parser.add_argument("--age-default", type=float, default=40.0, help="Fallback age when AGE is absent")
    parser.add_argument("--height-default", type=float, default=170.0, help="Fallback height when HEIGHT is absent")
    parser.add_argument("--no-anonymize", action="store_true", help="Copy DOCX files without anonymization; use only for controlled debugging")
    parser.add_argument("--doc-converter", choices=["auto", "libreoffice", "word"], default="auto", help="Converter for legacy .doc files. Default: auto = LibreOffice first, then Word/COM isolated subprocess")
    parser.add_argument("--conversion-timeout", type=int, default=60, help="Per-file conversion timeout in seconds for .doc files")
    args = parser.parse_args(argv)

    summary = build_validation_dataset(
        args.input,
        args.output,
        age_default=args.age_default,
        height_default=args.height_default,
        anonymize=not args.no_anonymize,
        doc_converter=args.doc_converter,
        conversion_timeout=args.conversion_timeout,
    )
    print("Phase-2 dataset completed")
    for key, value in summary.items():
        print(f"{key}: {value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
